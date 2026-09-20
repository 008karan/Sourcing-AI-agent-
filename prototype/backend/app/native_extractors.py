from __future__ import annotations

import email
import re
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document
from openpyxl import load_workbook

from .store import DATA_DIR, load_demo


def _money(text: str) -> float:
    return float(re.sub(r"[^0-9.]", "", str(text)))


def extract_packright(path: Path | None = None) -> dict[str, Any]:
    path = path or DATA_DIR / "vendor_a_packright_offer.xlsx"
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Commercial Offer"]
    facts = []
    for row_no, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if not row[1]:
            continue
        basis = "100_piece" if "100" in str(row[4]) else "piece"
        facts.append({
            "line_no": int(str(row[0]).split("-")[-1]),
            "sku": str(row[1]),
            "raw_basis": basis,
            "raw_price": float(row[6]),
            "raw_currency": "INR",
            "pack_size": int(row[5]),
            "lead_days": int(row[7]),
            "evidence": {"file": path.name, "sheet": "Commercial Offer", "cell_range": f"B{row_no}:I{row_no}", "source_text": " | ".join(str(x) for x in row)},
        })
    terms_ws = wb["Terms"]
    terms = {str(r[0]): r[1] for r in terms_ws.iter_rows(min_row=2, values_only=True) if r[0]}
    return {"vendor_id": "packright", "facts": facts, "terms": terms, "coverage": len(facts), "method": "native_xlsx"}


def extract_corrpro(path: Path | None = None) -> dict[str, Any]:
    path = path or DATA_DIR / "vendor_b_corrpro_quote.pdf"
    facts = []
    all_text = []
    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            all_text.append(text)
            for table in page.extract_tables():
                for row in table[1:]:
                    if not row or not row[0] or not str(row[0]).strip().isdigit():
                        continue
                    facts.append({
                        "line_no": int(row[0]),
                        "sku": row[1],
                        "raw_basis": "100_piece",
                        "raw_price": float(row[4]),
                        "raw_currency": "USD",
                        "pack_size": 100,
                        "lead_days": int(str(row[5]).replace("d", "")),
                        "evidence": {"file": path.name, "page": page_no, "source_text": " | ".join(str(x) for x in row)},
                    })
    text = "\n".join(all_text)
    rebate = None
    m = re.search(r"(\d+(?:\.\d+)?)% year-end volume rebate.*?USD\s*([\d,]+)", text, re.S | re.I)
    if m:
        rebate = {"percent": float(m.group(1)), "threshold_usd": float(m.group(2).replace(",", "")), "source_text": m.group(0)[:280]}
    return {
        "vendor_id": "corrpro",
        "facts": sorted(facts, key=lambda x: x["line_no"]),
        "terms": {"freight": "included", "currency": "USD", "rebate": rebate},
        "coverage": len(facts),
        "method": "native_pdf_table_plus_text",
    }


def extract_boxworks(path: Path | None = None) -> dict[str, Any]:
    path = path or DATA_DIR / "vendor_c_boxworks_offer.docx"
    doc = Document(path)
    facts = []
    table = doc.tables[0]
    for row_idx, row in enumerate(table.rows[1:], start=2):
        c = [x.text.strip() for x in row.cells]
        if not c[0].isdigit():
            continue
        facts.append({
            "line_no": int(c[0]),
            "sku": c[1],
            "raw_basis": "carton" if "carton" in c[3].lower() else "piece",
            "raw_price": _money(c[5]),
            "raw_currency": "INR",
            "pack_size": int(c[4]),
            "lead_days": int(re.search(r"\d+", c[6]).group()),
            "evidence": {"file": path.name, "table": 0, "row": row_idx, "source_text": " | ".join(c)},
        })
    narrative = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    match = re.search(r"₹([\d.]+) per piece", narrative)
    freight = float(match.group(1)) if match else None
    missing = [int(x) for x in re.findall(r"Lines?\s+([\d, ]+)\s+(?:are|is) not offered", narrative) for x in re.findall(r"\d+", x)]
    return {
        "vendor_id": "boxworks",
        "facts": facts,
        "terms": {"freight_per_piece_inr": freight, "missing_lines": missing, "quality_note": narrative[-480:]},
        "coverage": len(facts),
        "method": "native_docx_table_plus_text",
    }


def extract_greencarton(path: Path | None = None) -> dict[str, Any]:
    path = path or DATA_DIR / "vendor_e_greencarton_reply.eml"
    msg = email.message_from_bytes(path.read_bytes())
    body = msg.get_payload(decode=True).decode(msg.get_content_charset() or "utf-8", errors="replace")
    facts = []
    for line in body.splitlines():
        m = re.match(r"L(\d+)\s+(\S+):\s+(.*)$", line.strip())
        if not m:
            continue
        line_no, sku, rhs = int(m.group(1)), m.group(2), m.group(3)
        if rhs.lower().startswith("same as last year"):
            facts.append({
                "line_no": line_no, "sku": sku, "raw_basis": "same_as_last_year", "raw_price": None,
                "raw_currency": "INR", "pack_size": None, "lead_days": 16,
                "evidence": {"file": path.name, "source_text": line.strip()},
            })
        else:
            rate = re.search(r"₹\s*([\d.]+)\s*/\s*kg", rhs)
            if rate:
                facts.append({
                    "line_no": line_no, "sku": sku, "raw_basis": "kg", "raw_price": float(rate.group(1)),
                    "raw_currency": "INR", "pack_size": None, "lead_days": 16,
                    "evidence": {"file": path.name, "source_text": line.strip()},
                })
    return {
        "vendor_id": "greencarton",
        "facts": facts,
        "terms": {"freight_percent": 0.02, "gst": "extra", "lead_days": 16},
        "coverage": len(facts),
        "method": "native_email_regex",
    }


def extract_native(vendor_id: str, path: Path | None = None) -> dict[str, Any]:
    mapping = {
        "packright": extract_packright,
        "corrpro": extract_corrpro,
        "boxworks": extract_boxworks,
        "greencarton": extract_greencarton,
    }
    if vendor_id == "alphapack":
        return {"vendor_id": vendor_id, "facts": [], "terms": {}, "coverage": 0, "method": "vision_required"}
    if vendor_id not in mapping:
        raise KeyError(vendor_id)
    return mapping[vendor_id](path)
