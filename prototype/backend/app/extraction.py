from __future__ import annotations

import email
import json
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

from .ai import ai_extract_file, ai_extract_text, has_ai
from .native_extractors import extract_native
from .store import DATA_DIR, load_demo


def native_blocks(filename: str) -> dict[str, Any]:
    path = DATA_DIR / filename
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        wb = load_workbook(path, data_only=False, read_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    rows.append([v for v in row])
            sheets.append({"name": ws.title, "rows": rows})
        return {"type": "xlsx", "file": filename, "sheets": sheets}
    if suffix == ".docx":
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables = []
        for ti, table in enumerate(doc.tables):
            tables.append({"table_index": ti, "rows": [[cell.text for cell in row.cells] for row in table.rows]})
        return {"type": "docx", "file": filename, "paragraphs": paragraphs, "tables": tables}
    if suffix == ".eml":
        msg = email.message_from_bytes(path.read_bytes())
        body = ""
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    body += part.get_payload(decode=True).decode(part.get_content_charset() or "utf-8", errors="replace")
        else:
            body = msg.get_payload(decode=True).decode(msg.get_content_charset() or "utf-8", errors="replace")
        return {"type": "email", "file": filename, "subject": msg.get("Subject"), "from": msg.get("From"), "body": body}
    return {"type": suffix.lstrip("."), "file": filename, "size_bytes": path.stat().st_size}


def extraction_preview(vendor_id: str, run_ai: bool = False) -> dict[str, Any]:
    data = load_demo()
    vendor = data["vendors"][vendor_id]
    file_map = {
        "packright": "vendor_a_packright_offer.xlsx",
        "corrpro": "vendor_b_corrpro_quote.pdf",
        "boxworks": "vendor_c_boxworks_offer.docx",
        "alphapack": "vendor_d_alphapack_rate_card_photo.jpg",
        "greencarton": "vendor_e_greencarton_reply.eml",
    }
    filename = file_map[vendor_id]
    parsed = native_blocks(filename)
    structured = extract_native(vendor_id)
    result = {"vendor_id": vendor_id, "file": filename, "native": parsed, "structured": structured, "ai_mode": "not_run"}
    if run_ai:
        if not has_ai():
            result["ai_mode"] = "unavailable_no_key"
        else:
            rfx = data["rfx"]
            rfx_context = json.dumps({"event_id": rfx["event_id"], "items": [{"line_no": i["line_no"], "sku": i["sku"], "description": i["description"], "unit_weight_kg": i["unit_weight_kg"]} for i in rfx["items"]], "questionnaire": rfx["questionnaire"]}, ensure_ascii=False)
            if Path(filename).suffix.lower() in {".xlsx", ".docx", ".eml"}:
                result["ai"] = ai_extract_text(json.dumps(parsed, ensure_ascii=False), rfx_context)
            else:
                result["ai"] = ai_extract_file(filename, rfx_context)
            result["ai_mode"] = "assisted"
    return result
