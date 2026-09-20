"""Native-first extraction and deterministic, evidence-backed business projections.

No function in this module reads demo_state.json or normalized_truth.csv.
"""
from __future__ import annotations
import copy
import email
import hashlib
import json
import math
import re
from pathlib import Path
from uuid import uuid4
import pdfplumber
from docx import Document
from openpyxl import load_workbook
from . import repository as repo
from .ai import ai_extract_file, ai_extract_text, has_ai
from .native_extractors import extract_native
from .store import DATA_DIR

def source_blocks(path):
    suffix = path.suffix.lower()
    if suffix == '.xlsx':
        wb = load_workbook(path, data_only=True, read_only=True)
        blocks = [{'sheet':s.title, 'row':i, 'text':' | '.join(str(x) for x in row if x is not None)}
                  for s in wb for i,row in enumerate(s.values,1) if any(x is not None for x in row)]
        wb.close()
        return blocks
    if suffix == '.docx':
        doc = Document(path)
        return ([{'paragraph':i,'text':p.text} for i,p in enumerate(doc.paragraphs) if p.text] +
                [{'table':t,'row':i,'text':' | '.join(c.text for c in row.cells)}
                 for t,table in enumerate(doc.tables) for i,row in enumerate(table.rows,1)])
    if suffix == '.pdf':
        with pdfplumber.open(path) as pdf:
            return [{'page':i,'text':p.extract_text() or ''} for i,p in enumerate(pdf.pages,1)]
    if suffix in {'.eml','.txt'}:
        if suffix == '.txt':
            body = path.read_text()
        else:
            from email import policy
            msg = email.message_from_bytes(path.read_bytes(), policy=policy.default)
            body = '\n'.join(p.get_content() for p in msg.walk() if p.get_content_type()=='text/plain')
        return [{'body_span':[m.start(),m.end()],'text':m.group()} for m in re.finditer(r'[^\n]+',body)]
    return []

def terms_from_text(text):
    terms = {'source_text':text, 'freight':'unknown'}
    if re.search(r'freight[\s:|]*(?:is )?included|\bDDP\b', text, re.I):
        terms['freight']='included'
    flat = re.search(r'freight.{0,30}₹\s*([\d.]+)\s*per piece',text,re.I)
    pct = re.search(r'freight.{0,30}?([\d.]+)%',text,re.I)
    if flat:
        terms.update(freight='flat', freight_per_piece_inr=float(flat.group(1)))
    elif pct:
        terms.update(freight='percent',freight_percent=float(pct.group(1))/100)
    rebate = re.search(r'(\d+(?:\.\d+)?)% year-end volume rebate.*?USD\s*([\d,]+)',text,re.I|re.S)
    if rebate:
        terms['rebate']={'percent':float(rebate.group(1)),'threshold_usd':float(rebate.group(2).replace(',','')),'source_text':rebate.group()}
    return terms

def native_answers(blocks, rfx):
    answers = {}
    for b in blocks:
        if b.get('sheet') == 'Quality':
            for q in rfx['questionnaire']:
                if b['text'].startswith(q['question']+' | '):
                    answers[q['id']]={'answer':b['text'].split(' | ')[-1], 'confidence':.99,'evidence':b}
    text = '\n'.join(b['text'] for b in blocks if b.get('sheet')!='Commercial Offer')
    patterns = {
        'Q1':r'ISO 9001\s+(?:valid|yes|certified)',
        'Q2':r'(?:FSC CoC\s+(?:valid|yes)|hold FSC Chain-of-Custody certification)',
        'Q3':r'(?:BF specs yes|meet stated burst factor|confirm compliance with the minimum burst-factor requirements)',
        'Q6':r'lead time.{0,25}?(\d+)\s*(?:calendar )?days',
    }
    for qid,pattern in patterns.items():
        if qid in answers:
            continue
        for b in blocks:
            m = re.search(pattern,b['text'],re.I)
            if m:
                answers[qid]={'answer':m.group(1) if qid=='Q6' else 'yes','confidence':.99,'evidence':b}
                break
    return answers

def qualification(answers, text, rfx):
    pending = False
    for q in rfx['questionnaire']:
        if not q['mandatory']:
            continue
        a = answers.get(q['id'],{})
        answer = str(a.get('answer','')).lower().strip()
        if q['type']=='yes_no':
            if re.match(r'^(no|false)\b',answer):
                return 'fail'
            if not re.match(r'^(yes|true)\b',answer) or a.get('confidence',0)<.85:
                pending = True
        elif q['type']=='number':
            match = re.search(r'\d+(?:\.\d+)?',answer)
            if not match:
                pending = True
            elif ('maximum' in q and float(match.group())>q['maximum']) or ('minimum' in q and float(match.group())<q['minimum']):
                return 'fail'
    if re.search(r'renewal is in process|expires .*before', text,re.I):
        pending = True
    return 'pending' if pending else 'pass'

def context(rfx):
    return json.dumps({'items':[{k:v for k,v in i.items() if k!='baseline_unit_price_inr'} for i in rfx['items']], 'questionnaire':rfx['questionnaire']})

def extract(vendor_id, path, rfx, use_ai=True, native_fixture=False, progress=None):
    if progress:progress('parsing','Reading native structure and source locations')
    blocks = source_blocks(path)
    text = '\n'.join(b['text'] for b in blocks)
    result = None
    if native_fixture and path.suffix.lower() not in {'.jpg','.jpeg','.png','.webp'}:
        result = extract_native(vendor_id, path)
    if result is not None:
        facts = result['facts']
        answers = native_answers(blocks,rfx)
        method=result['method']
        unresolved=[]
        model=None
    elif use_ai and has_ai():
        if progress:progress('extracting','AI is interpreting source facts and ambiguity')
        ai = ai_extract_text(text,context(rfx)) if blocks and path.suffix.lower()!='.pdf' else ai_extract_file(path.name,context(rfx),path)
        facts=ai['quote_lines']
        answers={q['question_id']:q for q in ai['questionnaire_answers']}
        method='ai_vision' if not blocks else 'ai_native_blocks'
        unresolved=ai['unresolved']
        model='AI-assisted extraction'
        if not text:
            text='\n'.join(t['source_text'] for t in ai['commercial_terms'])
    else:
        return {'facts':[], 'terms':{}, 'questionnaire':{}, 'quality':'pending', 'method':'vision_required' if not blocks else 'ai_required', 'unresolved':['Runtime API key is required to interpret this document.'], 'blocks':blocks, 'model':None}
    terms=terms_from_text(text)
    if progress:progress('validating','Checking line mapping, currencies, units and qualification')
    seen=set()
    for fact in facts:
        line=fact.get('line_no')
        if not isinstance(line,int) or not 1<=line<=len(rfx['items']) or line in seen:
            raise ValueError('Extraction returned an invalid or duplicate RFx line. Review source mapping.')
        seen.add(line)
        if fact.get('sku') != rfx['items'][line-1]['sku']:
            fact['quote_status']='unclear'
            fact['mapping_issue']='Source SKU does not match this RFx line'
        fact.setdefault('quote_status','quoted')
        fact.setdefault('confidence',.99)
        fact.setdefault('evidence',{'file':path.name,'source_text':fact.get('source_text','')})
        fact['extraction_method']=method
    return {'facts':facts,'terms':terms,'questionnaire':answers,'quality':qualification(answers,text,rfx), 'method':method,'unresolved':unresolved,'blocks':blocks,'model':model}

def receive(vid,documents,message='',subject='Supplier response',sender_email=None):
    current=repo.read()
    if vid not in current['vendors']:raise KeyError('Unknown supplier')
    if current['rfx_status']!='approved':raise ValueError('Approve the RFx before receiving responses.')
    def mutate(d):
        v=d['vendors'][vid]
        v.update(documents=documents,message=message,subject=subject,response_status='received',received_at=repo.now())
        if sender_email:v['email']=sender_email
        # Old prices must not survive a newly received, not-yet-extracted offer.
        v.update(facts=[],terms={},questionnaire={},quality='pending',method='awaiting_processing',checksum=None)
        d['exceptions']=[e for e in d['exceptions'] if e['vendor_id']!=vid]
        d['stage']='collecting_responses'
        rebuild(d)
    return repo.change('response_received:'+vid,'supplier_inbox',mutate)

def receive_demo():
    for vid,(_,filename) in repo.VENDORS.items():
        path=DATA_DIR/filename
        receive(vid,[{'path':str(path),'filename':filename,'fixture':True}],subject='FY27 quote — demo supplier message')
    return repo.read()

def ingest_received(vid,use_ai=True,job_id=None):
    from .jobs import progress as notify
    progress=lambda stage,message:notify(job_id,vid,stage,message)
    d=repo.read();v=d['vendors'][vid];docs=v.get('documents',[])
    if not docs:
        progress('waiting','No supplier response received yet');return d
    combined_digest=hashlib.sha256(''.join(hashlib.sha256(Path(x['path']).read_bytes()).hexdigest() for x in docs).encode()).hexdigest()
    if v.get('packet_checksum')==combined_digest and v['response_status']=='processed' and v.get('method') not in {'vision_required','ai_required'}:
        progress('completed','Previously processed source is unchanged');return d
    progress('received',f'Received {len(docs)} source document(s)')
    results=[]
    for doc in docs:
        path=Path(doc['path'])
        fixture_name=repo.VENDORS.get(vid,(None,None))[1]
        fixture_path=DATA_DIR/fixture_name if fixture_name else None
        exact_fixture=bool(fixture_path) and path.suffix==fixture_path.suffix and hashlib.sha256(path.read_bytes()).digest()==hashlib.sha256(fixture_path.read_bytes()).digest()
        # Message-only cover notes add evidence/terms; the attached document supplies quotes.
        if doc.get('cover_note') and len(docs)>1:
            blocks=source_blocks(path);text='\n'.join(b['text'] for b in blocks)
            result={'facts':[],'terms':terms_from_text(text),'questionnaire':{},'quality':'pending','method':'native_email_cover','unresolved':[],'blocks':blocks,'model':None}
        else:result=extract(vid,path,d['rfx'],use_ai,native_fixture=exact_fixture,progress=progress)
        for f in result['facts']:
            f.setdefault('evidence',{})['source_path']=str(path)
            f['evidence']['file']=doc['filename']
        result['source_document']=doc
        results.append(result)
    main=next((r for r in results if r['facts']),results[0])
    merged=copy.deepcopy(main);facts={}
    for result in results:
        for f in result['facts']:
            if f['line_no'] in facts and any(f.get(k)!=facts[f['line_no']].get(k) for k in ['raw_price','raw_basis','raw_currency','pack_size']):
                facts[f['line_no']]['quote_status']='unclear'
                facts[f['line_no']]['mapping_issue']='Conflicting quotes across response attachments'
            else:facts[f['line_no']]=f
    merged['facts']=list(facts.values())
    alltext='\n'.join(r['terms'].get('source_text','') for r in results)
    merged['terms']=terms_from_text(alltext)
    merged['questionnaire']={k:v for r in results for k,v in r['questionnaire'].items()}
    merged['quality']=qualification(merged['questionnaire'],alltext,d['rfx'])
    merged['unresolved']=list(dict.fromkeys(x for r in results for x in r['unresolved']))
    progress('normalizing','Converting prices, freight and currency with explicit rules')
    result=commit_extraction(vid,Path(main['source_document']['path']),merged)
    result=repo.change('response_packet_processed:'+vid,'workflow',lambda d:d['vendors'][vid].update(packet_checksum=combined_digest))
    progress('completed',f"{len(merged['facts'])} line attempts · qualification {merged['quality']}")
    return result

def normalized(fact,terms,item,rfx):
    """All arithmetic is deterministic. Unknown inputs produce no landed price."""
    out={'unit_price':None,'landed_unit_cost':None,'transformation_steps':[],'blocking_exception':False}
    steps=out['transformation_steps']
    if fact.get('quote_status') in {'not_quoted','excluded'}:
        return out
    if fact.get('quote_status')=='unclear' or fact.get('confidence',0)<.85 or fact.get('mapping_issue'):
        return dict(out,blocking_exception=True,reason='Uncertain extraction or line mapping requires buyer review')
    if not fact.get('evidence_id'):
        return dict(out,blocking_exception=True,reason='Missing source evidence')
    raw=fact.get('raw_price'); basis=fact.get('raw_basis'); currency=fact.get('raw_currency')
    if currency not in {'INR','USD'}:
        return dict(out,blocking_exception=True,reason='Unsupported or missing currency')
    if basis=='same_as_last_year':
        if not fact.get('prior_year_approved'):
            return dict(out,blocking_exception=True,reason='Confirm the FY26 buyer baseline is the intended prior-year rate')
        unit=item['baseline_unit_price_inr']
        steps.append(f"Buyer-approved FY26 baseline: INR {unit}/piece; buyer_rfx_reference.xlsx")
    else:
        if not isinstance(raw,(int,float)) or not math.isfinite(raw) or raw<=0:
            return dict(out,blocking_exception=True,reason='Missing or invalid rate')
        if basis=='piece': unit=raw
        elif basis=='100_piece': unit=raw/100
        elif basis in {'carton','bundle'}:
            pack=fact.get('pack_size')
            if not isinstance(pack,(int,float)) or not math.isfinite(pack) or pack<=0:
                return dict(out,blocking_exception=True,reason='Missing or invalid pack size')
            unit=raw/pack
        elif basis=='kg': unit=raw*item['unit_weight_kg']
        else: return dict(out,blocking_exception=True,reason='Unsupported or missing unit basis')
        steps.append(f"{raw} {currency}/{basis}; " + (f"divide by {fact.get('pack_size')} pieces" if basis in {'carton','bundle'} else 'divide by 100' if basis=='100_piece' else f"multiply by buyer weight {item['unit_weight_kg']} kg/piece" if basis=='kg' else 'per-piece rate'))
    if currency=='USD':
        fx=rfx['commercial_rules']['fx_rate_usd_inr']; unit*=fx
        steps.append(f'USD × {fx} INR/USD — fixed RFx comparison assumption, not a live market rate')
    out['unit_price']=round(unit,2)
    if terms.get('freight')=='included':
        landed=unit; steps.append('Freight included per supplier source')
    elif terms.get('freight')=='flat':
        landed=unit+terms['freight_per_piece_inr']; steps.append(f"Add freight INR {terms['freight_per_piece_inr']}/piece")
    elif terms.get('freight')=='percent':
        landed=unit*(1+terms['freight_percent']); steps.append(f"Add freight {terms['freight_percent']*100:g}% of material value")
    else:
        return dict(out,blocking_exception=True,reason='Freight is unknown; landed cost withheld')
    steps.append('GST excluded; conditional rebates excluded')
    out['landed_unit_cost']=round(landed,2)
    return out

def rebuild(data):
    old={e['id']:e for e in data['exceptions']}
    exceptions=[]; rows=[]
    def issue(vid,line,kind,title,detail,eid=None,severity='high',blocking=False,confidence=.5):
        ident=f'{vid}:{line}:{kind}'
        previous=old.get(ident,{})
        exposure=(data['rfx']['items'][line-1]['annual_quantity']*data['rfx']['items'][line-1]['baseline_unit_price_inr']) if line else data['rfx']['expected_annual_spend_inr']
        e={'id':ident,'vendor_id':vid,'line_no':line,'type':kind,'title':title,'detail':detail,'evidence_id':eid,'severity':severity,'financial_exposure_inr':exposure,'priority_score':round(exposure*(1-confidence)*(1 if data['vendors'][vid]['quality']=='pass' else .2),2),'status':previous.get('status','blocking' if blocking else 'needs_review')}
        if previous.get('resolution'):e['resolution']=previous['resolution']
        exceptions.append(e)
    for vid,v in data['vendors'].items():
        if v['response_status']=='awaiting_ingestion':continue
        if v['quality']!='pass':
            issue(vid,0,'qualification','Mandatory qualification '+v['quality'], 'Supplier declarations fail a mandatory gate or require certification review. Excluded from qualified-only awards.',v.get('terms_evidence_id'),'critical',True)
        if v['terms'].get('rebate'):
            issue(vid,0,'conditional_discount','Conditional volume rebate',v['terms']['rebate']['source_text']+' — excluded from all award calculations.',v.get('terms_evidence_id'),'medium')
        if v.get('unresolved'):
            issue(vid,0,'document_review','Source interpretation notes','; '.join(v['unresolved']),v.get('terms_evidence_id'),'medium')
    for item in data['rfx']['items']:
        row={'line_no':item['line_no'],'sku':item['sku'],'description':item['description'],'qty':item['annual_quantity'],'baseline_unit_price_inr':item['baseline_unit_price_inr'],'vendors':{}}
        for vid,v in data['vendors'].items():
            fact=next((f for f in v['facts'] if f['line_no']==item['line_no']),None)
            if not fact:
                status='awaiting_ingestion' if v['response_status']=='awaiting_ingestion' else 'unclear' if v.get('method') in {'vision_required','ai_required'} else 'not_quoted'
                q={'status':status,'qualification':v['quality'],'landed_unit_cost':None,'unit_price':None,'confidence':0,'raw_basis':'','evidence_id':v.get('terms_evidence_id'),'blocking_exception':status=='unclear'}
                if status in {'not_quoted','unclear'}:
                    issue(vid,item['line_no'],'missing_quote' if status=='not_quoted' else 'extraction_required',f"Line {item['line_no']}: {'not quoted' if status=='not_quoted' else 'extraction needed'}",'No usable price was extracted; never treated as zero.',v.get('terms_evidence_id'))
            else:
                n=normalized(fact,v['terms'],item,data['rfx'])
                q={**copy.deepcopy(fact),**n,'status':fact.get('quote_status','quoted'),'qualification':v['quality']}
                if n['blocking_exception']:
                    issue(vid,item['line_no'],'review_price',f"Line {item['line_no']}: review price",n['reason'],fact['evidence_id'],'high',True,fact.get('confidence',.5))
                if fact.get('evidence_id'):
                    data['evidence'][fact['evidence_id']].update(line_no=item['line_no'],sku=item['sku'],description=item['description'],vendor_id=vid,vendor_name=v['name'],confidence=fact.get('confidence'),extraction_method=fact.get('extraction_method'),raw_fact=copy.deepcopy(fact),landed_unit_cost=n['landed_unit_cost'],normalized_unit_price=n['unit_price'],transformation_steps=n['transformation_steps'],buyer_override=fact.get('buyer_override'))
            row['vendors'][vid]=q
        rows.append(row)
    # Retain resolved exception records even after the condition disappears.
    present={e['id'] for e in exceptions}
    exceptions.extend(e for k,e in old.items() if k not in present and e.get('status')=='resolved_by_buyer')
    data['comparison']=rows
    data['exceptions']=sorted(exceptions,key=lambda e:(e['status']=='resolved_by_buyer',-e['priority_score']))

def commit_extraction(vid,path,result,expected=None):
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    def mutate(data):
        v=data['vendors'][vid]
        # A genuinely new response invalidates prior review decisions for this supplier.
        # Their prior facts and decisions remain in immutable older snapshots.
        data['exceptions']=[e for e in data['exceptions'] if e['vendor_id']!=vid]
        v.update({k:copy.deepcopy(result[k]) for k in ['facts','terms','questionnaire','quality','method','unresolved','model']})
        v.update(file=path.name,format=path.suffix[1:],source_path=str(path),checksum=digest,response_status='processed',received_at=repo.now())
        v['lead']=next((f.get('lead_days') for f in v['facts'] if f.get('lead_days')),None)
        eid=str(uuid4());v['terms_evidence_id']=eid
        data['evidence'][eid]={'id':eid,'type':v['format'],'file':path.name,'source_path':str(path),'source_text':v['terms'].get('source_text','') or '; '.join(v['unresolved']),'vendor_id':vid,'vendor_name':v['name'],'blocks':result.get('blocks',[]),'extraction_method':v['method'],'checksum':digest}
        for fact in v['facts']:
            eid=str(uuid4());fact['evidence_id']=eid;fact['fact_id']=str(uuid4())
            data['evidence'][eid]={'id':eid,'type':v['format'],'file':path.name,'source_path':str(path),**fact.get('evidence',{}),'source_text':fact.get('evidence',{}).get('source_text') or fact.get('source_text',''),'checksum':digest}
        rebuild(data)
        data['stage']='exception_review'
    return repo.change('ingest:'+vid,'extractor',mutate,expected)

def ingest(vid,use_ai=True,path=None):
    data=repo.read()
    if data['rfx_status']!='approved':raise ValueError('Approve the RFx before collecting responses.')
    v=data['vendors'][vid]
    fixture=path is None
    if path is None and vid not in repo.VENDORS:raise ValueError('This supplier has no bundled fixture; receive their response first.')
    path=Path(path) if path else DATA_DIR / repo.VENDORS[vid][1]
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    if v.get('checksum')==digest and v.get('method') not in {'vision_required','ai_required'}:
        return data  # preserve corrections on duplicate delivery/replayed graph node
    result=extract(vid,path,data['rfx'],use_ai,native_fixture=fixture)
    return commit_extraction(vid,path,result)

def resolve(eid,req):
    def mutate(data):
        exc=next((x for x in data['exceptions'] if x['id']==eid),None)
        if not exc:raise KeyError('Exception not found')
        if exc['status']=='resolved_by_buyer':raise ValueError('This exception is already resolved.')
        vid=exc['vendor_id'];v=data['vendors'][vid];line=exc['line_no'];action=req['action']
        reason=req.get('reason','').strip();actor=req.get('actor','').strip()
        if not reason or not actor:raise ValueError('Buyer name and reason are required.')
        fact=next((f for f in v['facts'] if f['line_no']==line),None)
        if line==0:
            if action!='accept':raise ValueError('Use acknowledge for supplier-level notes; qualification is not overridden.')
        else:
            if fact is None:
                fact={'line_no':line,'sku':data['rfx']['items'][line-1]['sku'],'raw_price':None,'raw_currency':'INR','raw_basis':'piece','pack_size':1,'lead_days':v['lead'],'confidence':0,'quote_status':'not_quoted','extraction_method':'buyer_entry'}
                v['facts'].append(fact)
            prior=copy.deepcopy(fact)
            if action=='accept':
                if fact.get('raw_basis')=='same_as_last_year':fact['prior_year_approved']=True
                elif fact.get('raw_price') is None:raise ValueError('An unreadable price cannot be accepted. Correct it with evidence or exclude it.')
                if fact.get('mapping_issue'):raise ValueError('Correct the line mapping before accepting this value.')
                fact.update(confidence=1,quote_status='quoted')
            elif action in {'correct','conversion'}:
                changes=req.get('changes',{})
                allowed={'raw_price','raw_currency','raw_basis','pack_size','lead_days'}
                if not changes or set(changes)-allowed:raise ValueError('Provide supported raw-value corrections.')
                fact.update(changes);fact.update(confidence=1,quote_status='quoted')
            elif action=='remap':
                target=req.get('target_line')
                if not isinstance(target,int) or not 1<=target<=len(data['rfx']['items']):raise ValueError('Invalid target line')
                existing=next((f for f in v['facts'] if f['line_no']==target),None)
                if existing and existing.get('quote_status')!='excluded':raise ValueError('Target line already has a fact; exclude it before remapping.')
                if existing:v['facts'].remove(existing)
                fact.update(line_no=target,sku=data['rfx']['items'][target-1]['sku'])
                fact.pop('mapping_issue',None)
            elif action in {'not_quoted','exclude'}:
                fact['quote_status']='not_quoted' if action=='not_quoted' else 'excluded'
            else:raise ValueError('Unsupported resolution')
            new_eid=str(uuid4());fact['evidence_id']=new_eid;fact['fact_id']=str(uuid4())
            fact['buyer_override']={'actor':actor,'reason':reason,'action':action,'timestamp':repo.now(),'previous_fact':prior,'parent_evidence_id':prior.get('evidence_id')}
            base=copy.deepcopy(data['evidence'].get(prior.get('evidence_id'),{}))
            data['evidence'][new_eid]={**base,'id':new_eid,'source_text':base.get('source_text','') or reason,'buyer_override':fact['buyer_override']}
            if req.get('reply'):
                data['evidence'][new_eid]['clarification_reply']=req['reply']
            if action in {'correct','conversion','accept'}:
                norm=normalized(fact,v['terms'],data['rfx']['items'][fact['line_no']-1],data['rfx'])
                if norm['blocking_exception']:raise ValueError(norm.get('reason','Correction is still invalid'))
        exc['status']='resolved_by_buyer';exc['resolution']={**req,'at':repo.now()}
        rebuild(data)
    return repo.change('resolve:'+eid,req.get('actor','buyer'),mutate,req.get('expected_version'))
